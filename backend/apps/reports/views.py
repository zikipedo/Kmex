"""Tableau de bord (§12) et rapports filtrables / exportables (§13)."""
import csv
import io
from datetime import date, datetime, timedelta
from decimal import Decimal

from django.db.models import Case, Count, DecimalField, ExpressionWrapper, F, Q, Sum, Value, When
from django.db.models.functions import Coalesce, TruncDate
from django.http import HttpResponse
from django.utils import timezone
from django.utils.dateparse import parse_date
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from rest_framework.decorators import api_view
from rest_framework.exceptions import NotFound, PermissionDenied
from rest_framework.response import Response

from apps.catalog.models import Product
from apps.core.models import AuditLog
from apps.core.services import audit
from apps.finance.models import CashMovement, Expense, Payment, RegisterSession, TreasuryAccount
from apps.inventory.models import StockLevel, StockMovement
from apps.purchasing.models import PurchaseInvoice
from apps.sales.models import CreditNote, Sale, SaleItem
from apps.sales.services import aging_buckets

DEC = DecimalField(max_digits=18, decimal_places=4)
ZERO = Value(Decimal("0"), output_field=DEC)
D0 = Decimal("0")


def period_bounds(request):
    p = request.query_params
    today = timezone.localdate()
    period = p.get("period", "month")
    if p.get("date_from") or p.get("date_to"):
        start = parse_date(p.get("date_from") or "") or today.replace(day=1)
        end = parse_date(p.get("date_to") or "") or today
    elif period == "today":
        start = end = today
    elif period == "week":
        start, end = today - timedelta(days=today.weekday()), today
    elif period == "year":
        start, end = today.replace(month=1, day=1), today
    elif period == "30d":
        start, end = today - timedelta(days=29), today
    else:
        start, end = today.replace(day=1), today
    span = (end - start).days + 1
    prev_end = start - timedelta(days=1)
    prev_start = prev_end - timedelta(days=span - 1)
    return start, end, prev_start, prev_end


def scoped_sales(request, start, end):
    u = request.user
    qs = Sale.objects.filter(company=u.company, issue_date__gte=start, issue_date__lte=end, warehouse_id__in=u.allowed_warehouse_ids())
    if request.query_params.get("warehouse"):
        qs = qs.filter(warehouse_id=request.query_params["warehouse"])
    if not u.has_code("reports.view"):
        qs = qs.filter(seller=u)  # le caissier ne voit que ses propres ventes (DASH-002)
    return qs


def sales_figures(sales_qs):
    agg = sales_qs.aggregate(n=Count("id"), ttc=Coalesce(Sum("total"), ZERO), ht=Coalesce(Sum("subtotal"), ZERO), disc=Coalesce(Sum("discount_total"), ZERO))
    credits = CreditNote.objects.filter(sale__in=sales_qs).aggregate(ttc=Coalesce(Sum("total"), ZERO), ht=Coalesce(Sum("subtotal"), ZERO))
    net_qty = ExpressionWrapper(F("quantity") - F("returned_qty"), output_field=DEC)
    items = SaleItem.objects.filter(sale__in=sales_qs).aggregate(
        rev=Coalesce(Sum(ExpressionWrapper(F("line_subtotal") * net_qty / F("quantity"), output_field=DEC)), ZERO),
        cogs=Coalesce(Sum(ExpressionWrapper(F("unit_cost") * net_qty, output_field=DEC)), ZERO),
    )
    n = agg["n"]
    return {
        "count": n,
        "revenue_ttc": agg["ttc"] - credits["ttc"],
        "revenue_ht": agg["ht"] - credits["ht"],
        "discounts": agg["disc"],
        "avg_basket": (agg["ttc"] / n) if n else D0,
        "cogs": items["cogs"],
        "margin": (items["rev"] - items["cogs"]).quantize(Decimal("0.0001")),
    }


def pct(cur, prev):
    if not prev:
        return None
    return round(float((Decimal(cur) - Decimal(prev)) / abs(Decimal(prev)) * 100), 1)


@api_view(["GET"])
def dashboard(request):
    u = request.user
    company = u.company
    start, end, pstart, pend = period_bounds(request)
    sales = scoped_sales(request, start, end)
    prev = scoped_sales(request, pstart, pend)
    cur_f, prev_f = sales_figures(sales), sales_figures(prev)
    can_profit = u.has_code("profit.view")
    can_cost = u.has_code("catalog.cost.view")
    can_fin = u.has_code("reports.finance") or u.has_code("cash.view")
    wh_ids = u.allowed_warehouse_ids()

    expenses = Expense.objects.filter(company=company, status="paid", expense_date__gte=start, expense_date__lte=end)
    if not u.has_code("expense.sensitive.view"):
        expenses = expenses.filter(category__is_sensitive=False)
    exp_total = expenses.aggregate(v=Coalesce(Sum("amount"), ZERO))["v"]

    levels = StockLevel.objects.filter(company=company, warehouse_id__in=wh_ids, product__status="active", product__track_stock=True)
    if request.query_params.get("warehouse"):
        levels = levels.filter(warehouse_id=request.query_params["warehouse"])
    per_product = levels.values("product_id", "product__min_stock").annotate(q=Sum("on_hand"))
    out_count = sum(1 for r in per_product if r["q"] <= 0)
    low_count = sum(1 for r in per_product if 0 < r["q"] <= (r["product__min_stock"] or 0))
    tracked = Product.objects.filter(company=company, status="active", track_stock=True).count()
    never = tracked - len(per_product)
    stock_value = levels.aggregate(v=Coalesce(Sum(F("on_hand") * F("avg_cost"), output_field=DEC), ZERO))["v"]

    open_sales = Sale.objects.filter(company=company, status__in=["issued", "partial"])
    receivables = open_sales.aggregate(v=Coalesce(Sum(F("total") - F("paid_amount") - F("credited_amount"), output_field=DEC), ZERO))["v"]
    overdue = open_sales.filter(due_date__lt=timezone.localdate()).aggregate(v=Coalesce(Sum(F("total") - F("paid_amount") - F("credited_amount"), output_field=DEC), ZERO))["v"]
    payables = PurchaseInvoice.objects.filter(company=company, status__in=["unpaid", "partial"]).aggregate(v=Coalesce(Sum(F("total") - F("paid_amount"), output_field=DEC), ZERO))["v"]

    # Séries journalières (période courante vs précédente)
    span = (end - start).days + 1
    daily = {r["d"]: r for r in sales.annotate(d=TruncDate("created_at")).values("d").annotate(v=Sum("total"), n=Count("id"))}
    daily_prev = {r["d"]: r["v"] for r in prev.annotate(d=TruncDate("created_at")).values("d").annotate(v=Sum("total"))}
    series = []
    for i in range(span if span <= 92 else 92):
        d = start + timedelta(days=i)
        pd = pstart + timedelta(days=i)
        series.append({"date": d.isoformat(), "revenue": str(daily.get(d, {}).get("v") or 0), "count": daily.get(d, {}).get("n") or 0, "previous": str(daily_prev.get(pd) or 0)})

    by_category = list(
        SaleItem.objects.filter(sale__in=sales).values("product__category__name", "product__category__color")
        .annotate(v=Sum("line_total"), q=Sum("quantity")).order_by("-v")[:8]
    )
    top_products = list(
        SaleItem.objects.filter(sale__in=sales).values("product_id", "product__name").annotate(q=Sum("quantity"), v=Sum("line_total")).order_by("-v")[:10]
    )
    by_method = list(
        Payment.objects.filter(company=company, direction="in", status="valid", paid_at__date__gte=start, paid_at__date__lte=end)
        .values("method__name", "method__color").annotate(v=Sum("amount"), n=Count("id")).order_by("-v")
    )
    exp_by_cat = list(expenses.values("category__name", "category__color").annotate(v=Sum("amount")).order_by("-v")[:6])
    recent = list(sales.select_related("customer").order_by("-created_at")[:6].values("id", "number", "customer__name", "total", "status", "created_at"))

    pending = {}
    if u.has_code("expense.approve"):
        pending["expenses"] = Expense.objects.filter(company=company, status="pending").exclude(created_by=u).count()
    if u.has_code("stock.adjust.approve"):
        from apps.inventory.models import StockAdjustment

        pending["adjustments"] = StockAdjustment.objects.filter(company=company, status="pending").exclude(created_by=u).count()
    if u.has_code("cash.session.validate"):
        pending["sessions"] = RegisterSession.objects.filter(company=company, status="closed").count()

    kpis = {
        "revenue": str(cur_f["revenue_ttc"]), "revenue_ht": str(cur_f["revenue_ht"]), "revenue_change": pct(cur_f["revenue_ttc"], prev_f["revenue_ttc"]),
        "sales_count": cur_f["count"], "sales_count_change": pct(cur_f["count"], prev_f["count"]),
        "avg_basket": str(cur_f["avg_basket"]), "avg_basket_change": pct(cur_f["avg_basket"], prev_f["avg_basket"]),
        "discounts": str(cur_f["discounts"]),
        "products_active": Product.objects.filter(company=company, status="active").count(),
        "out_of_stock": out_count + max(never, 0), "low_stock": low_count,
        "customers_new": sales.values("customer").distinct().count(),
        "expenses": str(exp_total),
    }
    if can_profit:
        kpis.update({
            "gross_margin": str(cur_f["margin"]), "gross_margin_change": pct(cur_f["margin"], prev_f["margin"]),
            "margin_rate": round(float(cur_f["margin"] / cur_f["revenue_ht"] * 100), 1) if cur_f["revenue_ht"] else None,
            "net_profit": str(cur_f["margin"] - exp_total),
        })
    if can_cost:
        kpis["stock_value"] = str(stock_value)
    if can_fin:
        kpis.update({"receivables": str(receivables), "receivables_overdue": str(overdue), "payables": str(payables)})

    treasury = []
    if can_fin:
        signed = Case(When(movements__direction="in", then=F("movements__amount")), default=-F("movements__amount"), output_field=DEC)
        treasury = [{"id": str(a.id), "name": a.name, "type": a.type, "balance": str(a.balance)}
                    for a in TreasuryAccount.objects.filter(company=company, is_active=True).annotate(balance=Coalesce(Sum(signed), ZERO))]

    session = None
    if u.has_code("cash.session"):
        from apps.finance.services import open_session_of, session_expected

        s = open_session_of(u)
        if s:
            ss = Sale.objects.filter(register_session=s).aggregate(n=Count("id"), v=Coalesce(Sum("total"), ZERO))
            session = {"id": str(s.id), "register": s.register.name, "opened_at": s.opened_at, "sales_count": ss["n"],
                       "sales_total": str(ss["v"]), "expected": {k: str(v) for k, v in session_expected(s).items()}}

    return Response({
        "period": {"start": start, "end": end, "previous_start": pstart, "previous_end": pend},
        "kpis": kpis,
        "series": series,
        "by_category": [{"name": r["product__category__name"], "color": r["product__category__color"], "value": str(r["v"]), "qty": str(r["q"])} for r in by_category],
        "top_products": [{"id": str(r["product_id"]), "name": r["product__name"], "qty": str(r["q"]), "value": str(r["v"])} for r in top_products],
        "payment_methods": [{"name": r["method__name"], "color": r["method__color"], "value": str(r["v"]), "count": r["n"]} for r in by_method],
        "expenses_by_category": [{"name": r["category__name"], "color": r["category__color"], "value": str(r["v"])} for r in exp_by_cat],
        "aging": {k: str(v) for k, v in aging_buckets(open_sales).items()} if can_fin else None,
        "recent_sales": [{**r, "id": str(r["id"]), "total": str(r["total"])} for r in recent],
        "pending": pending,
        "treasury": treasury,
        "session": session,
        "updated_at": timezone.now(),
    })


# --------------------------------------------------------------------------------------------
# Rapports
# --------------------------------------------------------------------------------------------

def col(key, label, type_="text", sensitive=None):
    return {"key": key, "label": label, "type": type_, "sensitive": sensitive}


def r_sales_by_day(request, start, end):
    qs = scoped_sales(request, start, end).annotate(d=TruncDate("created_at")).values("d").annotate(
        n=Count("id"), ht=Sum("subtotal"), tax=Sum("tax_total"), ttc=Sum("total"), disc=Sum("discount_total"), cost=Sum("cost_total")).order_by("d")
    rows = [{"date": r["d"], "count": r["n"], "discounts": r["disc"], "ht": r["ht"], "tax": r["tax"], "ttc": r["ttc"],
             "avg_basket": r["ttc"] / r["n"] if r["n"] else 0, "margin": r["ht"] - r["cost"]} for r in qs]
    return [col("date", "Date", "date"), col("count", "Ventes", "int"), col("discounts", "Remises", "money"), col("ht", "CA HT", "money"),
            col("tax", "Taxes", "money"), col("ttc", "CA TTC", "money"), col("avg_basket", "Panier moyen", "money"), col("margin", "Marge", "money", "profit.view")], rows


def r_sales_by_product(request, start, end):
    qs = SaleItem.objects.filter(sale__in=scoped_sales(request, start, end)).values("product__sku", "product__name", "product__category__name").annotate(
        q=Sum("quantity"), ret=Sum("returned_qty"), ht=Sum("line_subtotal"), ttc=Sum("line_total"),
        cost=Sum(F("unit_cost") * F("quantity"), output_field=DEC)).order_by("-ttc")
    total = sum((r["ttc"] for r in qs), D0) or 1
    rows = [{"sku": r["product__sku"], "name": r["product__name"], "category": r["product__category__name"], "qty": r["q"], "returned": r["ret"],
             "ttc": r["ttc"], "share": round(float(r["ttc"] / total * 100), 1), "margin": r["ht"] - r["cost"],
             "margin_rate": round(float((r["ht"] - r["cost"]) / r["ht"] * 100), 1) if r["ht"] else None} for r in qs]
    return [col("sku", "SKU"), col("name", "Produit"), col("category", "Catégorie"), col("qty", "Qté", "qty"), col("returned", "Retours", "qty"),
            col("ttc", "CA TTC", "money"), col("share", "Part %", "pct"), col("margin", "Marge", "money", "profit.view"),
            col("margin_rate", "Taux marge %", "pct", "profit.view")], rows


def r_sales_by_category(request, start, end):
    qs = SaleItem.objects.filter(sale__in=scoped_sales(request, start, end)).values("product__category__name").annotate(
        q=Sum("quantity"), ht=Sum("line_subtotal"), ttc=Sum("line_total"), cost=Sum(F("unit_cost") * F("quantity"), output_field=DEC)).order_by("-ttc")
    rows = [{"category": r["product__category__name"], "qty": r["q"], "ttc": r["ttc"], "margin": r["ht"] - r["cost"]} for r in qs]
    return [col("category", "Catégorie"), col("qty", "Qté", "qty"), col("ttc", "CA TTC", "money"), col("margin", "Marge", "money", "profit.view")], rows


def r_sales_by_seller(request, start, end):
    qs = scoped_sales(request, start, end).values("seller__full_name").annotate(
        n=Count("id"), ttc=Sum("total"), disc=Sum("discount_total"), cancelled=Count("id", filter=Q(status="cancelled"))).order_by("-ttc")
    rows = [{"seller": r["seller__full_name"] or "—", "count": r["n"], "ttc": r["ttc"], "discounts": r["disc"], "cancelled": r["cancelled"]} for r in qs]
    return [col("seller", "Vendeur / caissier"), col("count", "Ventes", "int"), col("ttc", "CA TTC", "money"), col("discounts", "Remises", "money"), col("cancelled", "Annulations", "int")], rows


def r_sales_by_customer(request, start, end):
    qs = scoped_sales(request, start, end).values("customer__code", "customer__name").annotate(
        n=Count("id"), ttc=Sum("total"), due=Sum(F("total") - F("paid_amount") - F("credited_amount"), output_field=DEC)).order_by("-ttc")
    rows = [{"code": r["customer__code"], "customer": r["customer__name"], "count": r["n"], "ttc": r["ttc"], "due": r["due"]} for r in qs]
    return [col("code", "Code"), col("customer", "Client"), col("count", "Achats", "int"), col("ttc", "CA TTC", "money"), col("due", "Reste dû", "money")], rows


def r_sales_journal(request, start, end):
    qs = scoped_sales(request, start, end).select_related("customer", "seller").order_by("created_at")
    rows = [{"number": s.number, "date": s.created_at, "customer": s.customer.name, "seller": s.seller.full_name if s.seller else "",
             "ht": s.subtotal, "tax": s.tax_total, "ttc": s.total, "paid": s.paid_amount, "status": s.get_status_display()} for s in qs]
    return [col("number", "N°"), col("date", "Date", "datetime"), col("customer", "Client"), col("seller", "Vendeur"), col("ht", "HT", "money"),
            col("tax", "Taxes", "money"), col("ttc", "TTC", "money"), col("paid", "Payé", "money"), col("status", "Statut")], rows


def r_payments(request, start, end):
    qs = Payment.objects.filter(company=request.user.company, direction="in", status="valid", paid_at__date__gte=start, paid_at__date__lte=end).values(
        "method__name").annotate(n=Count("id"), v=Sum("amount")).order_by("-v")
    rows = [{"method": r["method__name"], "count": r["n"], "amount": r["v"]} for r in qs]
    return [col("method", "Moyen de paiement"), col("count", "Opérations", "int"), col("amount", "Montant", "money")], rows


def r_stock_valuation(request, start, end):
    u = request.user
    qs = StockLevel.objects.filter(company=u.company, warehouse_id__in=u.allowed_warehouse_ids()).exclude(product__status="archived").select_related("product", "product__category", "warehouse")
    if request.query_params.get("warehouse"):
        qs = qs.filter(warehouse_id=request.query_params["warehouse"])
    rows = [{"sku": l.product.sku, "name": l.product.name, "category": l.product.category.name, "warehouse": l.warehouse.name, "qty": l.on_hand,
             "avg_cost": l.avg_cost, "value": l.on_hand * l.avg_cost, "retail_value": l.on_hand * l.product.price_retail} for l in qs.order_by("product__name")]
    return [col("sku", "SKU"), col("name", "Produit"), col("category", "Catégorie"), col("warehouse", "Dépôt"), col("qty", "Quantité", "qty"),
            col("avg_cost", "CMUP", "money", "catalog.cost.view"), col("value", "Valeur (coût)", "money", "catalog.cost.view"), col("retail_value", "Valeur (vente)", "money")], rows


def r_low_stock(request, start, end):
    u = request.user
    qs = StockLevel.objects.filter(company=u.company, warehouse_id__in=u.allowed_warehouse_ids(), product__status="active", product__track_stock=True).filter(
        on_hand__lte=F("product__min_stock")).select_related("product", "warehouse", "product__main_supplier")
    since = timezone.localdate() - timedelta(days=30)
    rows = []
    for l in qs.order_by("on_hand"):
        sold = SaleItem.objects.filter(product=l.product, sale__warehouse=l.warehouse, sale__issue_date__gte=since).aggregate(v=Coalesce(Sum("quantity"), ZERO))["v"]
        rows.append({"sku": l.product.sku, "name": l.product.name, "warehouse": l.warehouse.name, "qty": l.on_hand, "min": l.product.min_stock,
                     "avg_daily": round(sold / 30, 2), "suggested": max((l.product.max_stock or l.product.min_stock * 3) - l.on_hand, 0),
                     "supplier": l.product.main_supplier.name if l.product.main_supplier else "", "state": "Rupture" if l.on_hand <= 0 else "Faible"})
    return [col("sku", "SKU"), col("name", "Produit"), col("warehouse", "Dépôt"), col("qty", "Stock", "qty"), col("min", "Minimum", "qty"),
            col("avg_daily", "Ventes / jour", "qty"), col("suggested", "Qté suggérée", "qty"), col("supplier", "Fournisseur"), col("state", "État")], rows


def r_movements(request, start, end):
    u = request.user
    qs = StockMovement.objects.filter(company=u.company, warehouse_id__in=u.allowed_warehouse_ids(), created_at__date__gte=start, created_at__date__lte=end).select_related("product", "warehouse", "user").order_by("-id")[:5000]
    rows = [{"date": m.created_at, "product": m.product.name, "warehouse": m.warehouse.name, "type": m.get_movement_type_display(), "qty": m.quantity,
             "before": m.qty_before, "after": m.qty_after, "document": m.document_number, "user": m.user.full_name if m.user else "", "reason": m.reason} for m in qs]
    return [col("date", "Date", "datetime"), col("product", "Produit"), col("warehouse", "Dépôt"), col("type", "Type"), col("qty", "Qté", "qty"),
            col("before", "Avant", "qty"), col("after", "Après", "qty"), col("document", "Document"), col("user", "Utilisateur"), col("reason", "Motif")], rows


def r_dormant(request, start, end):
    u = request.user
    days = int(u.company.setting("dormant_days") or 90)
    since = timezone.localdate() - timedelta(days=days)
    sold_ids = SaleItem.objects.filter(sale__company=u.company, sale__issue_date__gte=since).values_list("product_id", flat=True).distinct()
    qs = StockLevel.objects.filter(company=u.company, on_hand__gt=0, warehouse_id__in=u.allowed_warehouse_ids()).exclude(product_id__in=sold_ids).select_related("product", "warehouse")
    rows = [{"sku": l.product.sku, "name": l.product.name, "warehouse": l.warehouse.name, "qty": l.on_hand, "value": l.on_hand * l.avg_cost} for l in qs]
    return [col("sku", "SKU"), col("name", "Produit"), col("warehouse", "Dépôt"), col("qty", "Stock", "qty"), col("value", "Valeur immobilisée", "money", "catalog.cost.view")], rows


def r_receivables(request, start, end):
    today = timezone.localdate()
    rows = {}
    for s in Sale.objects.filter(company=request.user.company, status__in=["issued", "partial"]).select_related("customer"):
        r = rows.setdefault(s.customer_id, {"customer": s.customer.name, "phone": s.customer.phone, "0_30": D0, "31_60": D0, "61_90": D0, "90_plus": D0, "total": D0})
        age = (today - (s.due_date or s.issue_date)).days
        k = "0_30" if age <= 30 else "31_60" if age <= 60 else "61_90" if age <= 90 else "90_plus"
        due = s.balance
        r[k] += due
        r["total"] += due
    data = sorted(rows.values(), key=lambda r: -r["total"])
    return [col("customer", "Client"), col("phone", "Téléphone"), col("0_30", "0-30 j", "money"), col("31_60", "31-60 j", "money"),
            col("61_90", "61-90 j", "money"), col("90_plus", "+90 j", "money"), col("total", "Total dû", "money")], data


def r_payables(request, start, end):
    today = timezone.localdate()
    rows = [{"number": i.number, "supplier": i.supplier.name, "ref": i.supplier_ref, "issue": i.issue_date, "due": i.due_date, "total": i.total,
             "paid": i.paid_amount, "balance": i.balance, "late_days": max((today - i.due_date).days, 0)}
            for i in PurchaseInvoice.objects.filter(company=request.user.company, status__in=["unpaid", "partial"]).select_related("supplier").order_by("due_date")]
    return [col("number", "N°"), col("supplier", "Fournisseur"), col("ref", "Réf. fournisseur"), col("issue", "Date", "date"), col("due", "Échéance", "date"),
            col("total", "Total", "money"), col("paid", "Payé", "money"), col("balance", "Reste dû", "money"), col("late_days", "Retard (j)", "int")], rows


def _expenses(request, start, end):
    qs = Expense.objects.filter(company=request.user.company, expense_date__gte=start, expense_date__lte=end)
    if not request.user.has_code("expense.sensitive.view"):
        qs = qs.filter(category__is_sensitive=False)
    return qs


def r_expenses_by_category(request, start, end):
    qs = _expenses(request, start, end).filter(status="paid").values("category__name", "category__monthly_budget").annotate(v=Sum("amount"), n=Count("id")).order_by("-v")
    total = sum((r["v"] for r in qs), D0) or 1
    rows = [{"category": r["category__name"], "count": r["n"], "amount": r["v"], "share": round(float(r["v"] / total * 100), 1), "budget": r["category__monthly_budget"]} for r in qs]
    return [col("category", "Catégorie"), col("count", "Nombre", "int"), col("amount", "Montant", "money"), col("share", "Part %", "pct"), col("budget", "Budget mensuel", "money")], rows


def r_expenses_journal(request, start, end):
    qs = _expenses(request, start, end).select_related("category", "created_by", "warehouse", "method").order_by("expense_date")
    rows = [{"number": e.number, "date": e.expense_date, "category": e.category.name, "description": e.description, "payee": e.payee_name,
             "warehouse": e.warehouse.name if e.warehouse else "Siège", "method": e.method.name if e.method else "", "amount": e.amount,
             "status": e.get_status_display(), "author": e.created_by.full_name if e.created_by else "", "receipts": e.attachments.count()} for e in qs]
    return [col("number", "N°"), col("date", "Date", "date"), col("category", "Catégorie"), col("description", "Motif"), col("payee", "Bénéficiaire"),
            col("warehouse", "Dépôt"), col("method", "Moyen"), col("amount", "Montant", "money"), col("status", "Statut"), col("author", "Saisi par"), col("receipts", "Justif.", "int")], rows


def r_profit(request, start, end):
    """Résultat de gestion simplifié : CA − COGS − dépenses = bénéfice net estimé (RG-DEP-11)."""
    f = sales_figures(scoped_sales(request, start, end))
    exp = _expenses(request, start, end).filter(status="paid").aggregate(v=Coalesce(Sum("amount"), ZERO))["v"]
    rows = [
        {"label": "Chiffre d'affaires HT (net des avoirs)", "amount": f["revenue_ht"]},
        {"label": "Coût des marchandises vendues (COGS)", "amount": -f["cogs"]},
        {"label": "Marge brute", "amount": f["margin"]},
        {"label": "Dépenses de fonctionnement payées", "amount": -exp},
        {"label": "Bénéfice net estimé", "amount": f["margin"] - exp},
    ]
    return [col("label", "Poste"), col("amount", "Montant", "money")], rows


def r_cash_sessions(request, start, end):
    qs = RegisterSession.objects.filter(company=request.user.company, opened_at__date__gte=start, opened_at__date__lte=end).select_related("register", "opened_by").order_by("-opened_at")
    rows = [{"z": s.z_number or "—", "register": s.register.name, "cashier": s.opened_by.full_name, "opened": s.opened_at, "closed": s.closed_at,
             "difference": s.difference, "reason": s.difference_reason, "status": s.get_status_display()} for s in qs]
    return [col("z", "Rapport Z"), col("register", "Caisse"), col("cashier", "Caissier"), col("opened", "Ouverture", "datetime"), col("closed", "Clôture", "datetime"),
            col("difference", "Écart", "money"), col("reason", "Justification"), col("status", "Statut")], rows


def r_taxes(request, start, end):
    qs = SaleItem.objects.filter(sale__in=scoped_sales(request, start, end)).values("tax_rate").annotate(base=Sum("line_subtotal"), tax=Sum("tax_amount")).order_by("tax_rate")
    rows = [{"rate": r["tax_rate"], "base": r["base"], "tax": r["tax"]} for r in qs]
    return [col("rate", "Taux %", "pct"), col("base", "Base HT", "money"), col("tax", "Taxe collectée", "money")], rows


def r_user_activity(request, start, end):
    qs = AuditLog.objects.filter(company=request.user.company, created_at__date__gte=start, created_at__date__lte=end).values("user_name").annotate(
        n=Count("id"), sales=Count("id", filter=Q(entity_type="sale", action="VALIDATE")), cancels=Count("id", filter=Q(action="CANCEL")),
        logins=Count("id", filter=Q(action="LOGIN"))).order_by("-n")
    rows = [{"user": r["user_name"], "actions": r["n"], "sales": r["sales"], "cancels": r["cancels"], "logins": r["logins"]} for r in qs]
    return [col("user", "Utilisateur"), col("actions", "Actions", "int"), col("sales", "Ventes validées", "int"), col("cancels", "Annulations", "int"), col("logins", "Connexions", "int")], rows


REPORTS = {
    "sales-by-day": ("Ventes par jour", "reports.view", r_sales_by_day),
    "sales-by-product": ("Ventes par produit", "reports.view", r_sales_by_product),
    "sales-by-category": ("Ventes par catégorie", "reports.view", r_sales_by_category),
    "sales-by-seller": ("Ventes par vendeur", "reports.view", r_sales_by_seller),
    "sales-by-customer": ("Ventes par client", "reports.view", r_sales_by_customer),
    "sales-journal": ("Journal des ventes", "reports.view", r_sales_journal),
    "payments": ("Encaissements par moyen", "reports.finance", r_payments),
    "stock-valuation": ("Valeur du stock", "reports.view", r_stock_valuation),
    "low-stock": ("Stock faible et ruptures", "reports.view", r_low_stock),
    "movements": ("Mouvements de stock", "reports.view", r_movements),
    "dormant": ("Produits dormants", "reports.view", r_dormant),
    "receivables": ("Créances clients (balance âgée)", "reports.finance", r_receivables),
    "payables": ("Dettes fournisseurs", "reports.finance", r_payables),
    "expenses-by-category": ("Dépenses par catégorie", ["reports.finance", "expense.view"], r_expenses_by_category),
    "expenses-journal": ("Journal des dépenses", ["reports.finance", "expense.view"], r_expenses_journal),
    "profit": ("Résultat de gestion", "profit.view", r_profit),
    "cash-sessions": ("Écarts de caisse", "reports.finance", r_cash_sessions),
    "taxes": ("Taxes collectées", "reports.finance", r_taxes),
    "user-activity": ("Activité des utilisateurs", "audit.view", r_user_activity),
}


def _safe(v):
    """Protection contre l'injection de formules CSV/Excel (§23.1)."""
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@"):
        return "'" + v
    return v


def _plain(v):
    if isinstance(v, Decimal):
        return float(v)
    if isinstance(v, datetime):
        return (timezone.localtime(v) if timezone.is_aware(v) else v).strftime("%d/%m/%Y %H:%M")
    if isinstance(v, date):
        return v.strftime("%d/%m/%Y")
    return v


@api_view(["GET"])
def report(request, name):
    if name not in REPORTS:
        raise NotFound("Rapport inconnu.")
    title, perm, fn = REPORTS[name]
    perms = perm if isinstance(perm, list) else [perm]
    if not any(request.user.has_code(p) for p in perms):
        raise PermissionDenied("Vous n'avez pas accès à ce rapport.")
    start, end, _, _ = period_bounds(request)
    columns, rows = fn(request, start, end)
    columns = [c for c in columns if not c["sensitive"] or request.user.has_code(c["sensitive"])]
    keys = [c["key"] for c in columns]
    totals = {}
    for c in columns if name != "profit" else []:
        if c["type"] in ("money", "int", "qty") and c["key"] not in ("avg_cost", "min", "avg_daily", "late_days"):
            totals[c["key"]] = sum((Decimal(str(r.get(c["key"]) or 0)) for r in rows), D0)

    fmt = request.query_params.get("export")
    if fmt in ("csv", "xlsx"):
        if not request.user.has_code("export") and not request.user.has_code("expense.export"):
            raise PermissionDenied("Vous n'avez pas la permission d'exporter.")
        audit(request.user.company, request.user, "EXPORT", "report", label=title, new={"format": fmt, "rows": len(rows), "from": start, "to": end})
        filename = f"stockpro-{name}-{start}-{end}.{fmt}"
        if fmt == "csv":
            buf = io.StringIO()
            buf.write("﻿")
            w = csv.writer(buf, delimiter=";")
            w.writerow([c["label"] for c in columns])
            for r in rows:
                w.writerow([_safe(_plain(r.get(k))) for k in keys])
            resp = HttpResponse(buf.getvalue(), content_type="text/csv; charset=utf-8")
        else:
            wb = Workbook()
            ws = wb.active
            ws.title = title[:30]
            ws.append([c["label"] for c in columns])
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="7C3AED")
            for r in rows:
                ws.append([_safe(_plain(r.get(k))) for k in keys])
            for i, c in enumerate(columns, start=1):
                ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = max(12, len(c["label"]) + 4)
            out = io.BytesIO()
            wb.save(out)
            resp = HttpResponse(out.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        resp["Content-Disposition"] = f'attachment; filename="{filename}"'
        return resp

    def ser(v):
        if isinstance(v, Decimal):
            return str(v)
        return v

    return Response({
        "name": name, "title": title, "period": {"start": start, "end": end}, "columns": columns,
        "rows": [{k: ser(r.get(k)) for k in keys} for r in rows], "totals": {k: str(v) for k, v in totals.items()},
    })


@api_view(["GET"])
def report_list(request):
    out = []
    for key, (title, perm, _fn) in REPORTS.items():
        perms = perm if isinstance(perm, list) else [perm]
        if any(request.user.has_code(p) for p in perms):
            out.append({"key": key, "title": title})
    return Response(out)
